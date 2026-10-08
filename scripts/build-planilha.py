#!/usr/bin/env python3
"""Gera a planilha de inscrições do Behold Movement (.xlsx para subir no Google Drive).

Abas:
  Inscrições  uma linha por pessoa, preenchida pelo n8n (chave: Telefone)
  Resumo      totais por ministério, cidade, bairro e status de envio (fórmulas)

Uso: python3 scripts/build-planilha.py [saida.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BLUE = "1F5BF2"
INK = "16171C"
PAPER = "ECEAE5"
FONT = "Arial"

# Os nomes das colunas precisam ser iguais aos do workflow do n8n.
COLUNAS = [
    ("ID", 16),
    ("Data/hora", 20),
    ("Nome", 28),
    ("Telefone", 17),
    ("E-mail", 30),
    ("Endereço", 32),
    ("Bairro", 22),
    ("Cidade", 20),
    ("Igreja", 26),
    ("Ministério", 18),
    ("Consentimento", 15),
    ("Status e-mail", 16),
    ("Status WhatsApp", 17),
    ("Conversa Chatwoot", 18),
    ("Observações", 36),
]
COL = {nome: get_column_letter(i + 1) for i, (nome, _) in enumerate(COLUNAS)}

MINISTERIOS = ["Evangelizando", "Liderando grupos", "Mídia", "Louvor",
               "Suprimento", "Ação social", "Intercessão"]
CIDADES = ["Taboão da Serra", "Embu das Artes", "Itapecerica da Serra",
           "São Paulo", "Outra cidade"]
BAIRROS = ["Jardim Saporito", "Jardim São Judas", "Jardim Comunitário",
           "Parque Pinheiros", "Jardim Salete", "Saint Moritz",
           "Jardim Record", "Jardim Clementino"]

head_font = Font(name=FONT, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=BLUE)
title_font = Font(name=FONT, bold=True, size=14, color=INK)
section_font = Font(name=FONT, bold=True, color=BLUE)
body_font = Font(name=FONT, color=INK)
label_fill = PatternFill("solid", fgColor=PAPER)


def rng(nome: str) -> str:
    c = COL[nome]
    return f"'Inscrições'!{c}:{c}"


def build(out: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Inscrições"
    for i, (nome, largura) in enumerate(COLUNAS, start=1):
        cell = ws.cell(row=1, column=i, value=nome)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}1"
    # O n8n grava em modo RAW, então +5511... fica como texto, sem virar número.

    rs = wb.create_sheet("Resumo")
    rs.column_dimensions["A"].width = 30
    rs.column_dimensions["B"].width = 14
    rs["A1"] = "Behold Movement · Inscrições Dia 31"
    rs["A1"].font = title_font

    row = 3

    def secao(titulo: str, linhas: list[tuple[str, str]]) -> None:
        nonlocal row
        rs.cell(row=row, column=1, value=titulo).font = section_font
        rs.cell(row=row, column=2, value="Pessoas").font = section_font
        row += 1
        for rotulo, formula in linhas:
            a = rs.cell(row=row, column=1, value=rotulo)
            a.font = body_font
            a.fill = label_fill
            b = rs.cell(row=row, column=2, value=formula)
            b.font = body_font
            row += 1
        row += 1

    # Colunas inteiras (C:C) funcionam no Excel e no Sheets; o -1 desconta o cabeçalho.
    total = f"=COUNTA({rng('Nome')})-1"
    secao("Geral", [
        ("Total de inscrições", total),
        ("E-mails enviados", f'=COUNTIF({rng("Status e-mail")},"enviado")'),
        ("WhatsApp enviados", f'=COUNTIF({rng("Status WhatsApp")},"enviado")'),
        ("Envios com erro",
         f'=COUNTIF({rng("Status e-mail")},"erro*")+COUNTIF({rng("Status WhatsApp")},"erro*")'),
    ])

    inicio = row + 1
    secao("Por ministério", [(m, f'=COUNTIF({rng("Ministério")},A{inicio + i})')
                              for i, m in enumerate(MINISTERIOS)])

    inicio = row + 1
    secao("Por cidade", [(c, f'=COUNTIF({rng("Cidade")},A{inicio + i})')
                          for i, c in enumerate(CIDADES)])

    inicio = row + 1
    linhas = [(b, f'=COUNTIF({rng("Bairro")},A{inicio + i})') for i, b in enumerate(BAIRROS)]
    fim = inicio + len(BAIRROS) - 1
    linhas.append(("Outros bairros", f"=COUNTA({rng('Nome')})-1-SUM(B{inicio}:B{fim})"))
    secao("Bairros prioritários", linhas)

    rs.cell(row=row, column=1,
            value="Os números são calculados a partir da aba Inscrições, que o n8n preenche a cada inscrição no site.").font = Font(name=FONT, italic=True, color="5D616C")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"ok: {out}")


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Behold-Inscricoes.xlsx"))
